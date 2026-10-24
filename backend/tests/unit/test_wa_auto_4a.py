"""WA-AUTO-4a: knowledge sources, chunking, retrieval, and the cache split."""
from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app.services import knowledge_service as ks
from tests.unit.test_account_plans_1 import FakeDB, _Q

ORG = "00000000-0000-0000-0000-00000000aa01"
OTHER = "00000000-0000-0000-0000-00000000aa02"


class _KQ(_Q):
    def delete(self):
        self.op = "delete"
        return self

    def maybe_single(self):
        return self

    def execute(self):
        if self.op == "delete":
            rows = self.db.tables.setdefault(self.name, [])
            gone = [r for r in rows if self._match(r)]
            self.db.tables[self.name] = [r for r in rows if r not in gone]
            return SimpleNamespace(data=gone)
        return super().execute()


class KDB(FakeDB):
    """FakeDB plus delete and a stand-in for the knowledge_search function (matches words in active sources)."""
    def __init__(self, **tables):
        tables.setdefault("knowledge_sources", [])
        tables.setdefault("knowledge_chunks", [])
        tables.setdefault("audit_logs", [])
        super().__init__(**tables)
        self.rpc_error = False

    def table(self, name):
        return _KQ(self, name)

    def rpc(self, fn, params):
        db = self

        class _R:
            def execute(_s):
                if db.rpc_error:
                    raise RuntimeError("down")
                words = [w.rstrip(":*") for w in params["p_query"].split(" | ")]
                active = {s["id"]: s for s in db.tables["knowledge_sources"] if s["org_id"] == params["p_org"] and s["status"] == "active"}
                hits = [{"id": c["id"], "source_id": c["source_id"], "title": active[c["source_id"]]["title"], "content": c["content"], "rank": 1.0}
                        for c in db.tables["knowledge_chunks"] if c["source_id"] in active and any(w in c["content"].lower() for w in words)]
                return SimpleNamespace(data=hits[: params["p_limit"]])
        return _R()


# -- chunking ------------------------------------------------------------------------------------------------------
def test_chunks_are_small_and_keep_all_the_text():
    text = "\n\n".join(f"Paragraph {i}. " + ("word " * 60) for i in range(40))
    chunks = ks.chunk_text(text)
    assert len(chunks) > 3 and all(len(c) <= ks.MAX_CHUNK_CHARS for c in chunks)
    assert "".join(c.replace("\n", "").replace(" ", "") for c in chunks) == text.replace("\n", "").replace(" ", "")


def test_a_very_long_paragraph_is_cut_at_sentences():
    para = " ".join(f"Sentence number {i} is here." for i in range(300))
    chunks = ks.chunk_text(para)
    assert len(chunks) > 3 and all(len(c) <= ks.MAX_CHUNK_CHARS for c in chunks)
    assert all(c.endswith(".") for c in chunks)


def test_unbroken_text_is_still_cut():
    chunks = ks.chunk_text("x" * 5000)
    assert all(len(c) <= ks.MAX_CHUNK_CHARS for c in chunks) and sum(len(c) for c in chunks) == 5000


def test_clean_text_removes_junk():
    assert ks.clean_text("a\x00b\r\n\r\n\r\n\r\nc   d") == "a b\n\nc d"


# -- reading files -----------------------------------------------------------------------------------------------
def test_text_and_markdown_files_are_read():
    assert "hello" in ks.extract_text("a.txt", b"hello there")
    assert "hi" in ks.extract_text("a.MD", b"# hi")


def test_docx_is_read_including_tables():
    from docx import Document
    d = Document()
    d.add_paragraph("Delivery takes two days.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Lagos", "N2,000"
    buf = io.BytesIO()
    d.save(buf)
    out = ks.extract_text("price.docx", buf.getvalue())
    assert "Delivery takes two days." in out and "Lagos | N2,000" in out


def test_pdf_is_read():
    from pypdf import PdfWriter
    from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
    w = PdfWriter()
    page = w.add_blank_page(width=300, height=300)
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 150 Td (Opening hours nine to five) Tj ET")
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    page[NameObject("/Contents")] = w._add_object(stream)
    buf = io.BytesIO()
    w.write(buf)
    assert "Opening hours nine to five" in ks.extract_text("hours.pdf", buf.getvalue())


@pytest.mark.parametrize("name,data,code", [("a.exe", b"x", "UNSUPPORTED_FILE"), ("a.txt", b"", "EMPTY_FILE"),
                                            ("a.pdf", b"not a pdf", "UNREADABLE_FILE"),
                                            ("a.txt", b"x" * (ks.MAX_FILE_BYTES + 1), "FILE_TOO_LARGE")])
def test_bad_files_are_refused_with_a_plain_message(name, data, code):
    with pytest.raises(HTTPException) as e:
        ks.extract_text(name, data)
    assert e.value.detail["code"] == code


# -- sources: draft, preview, start, remove ------------------------------------------------------------------
def test_added_text_is_a_draft_until_started():
    db = KDB()
    out = ks.add_text(db, ORG, "u1", "Delivery", "We deliver in Lagos within two days. " * 5)
    assert out["status"] == "draft" and out["chunk_count"] == 1 and out["samples"]
    assert ks.has_active(db, ORG) is False
    assert ks.retrieve(db, ORG, "delivery lagos") == []                 # a draft is never given to the bot
    ks.activate(db, ORG, out["id"])
    assert ks.has_active(db, ORG) is True
    assert ks.retrieve(db, ORG, "do you deliver in lagos")[0]["chunk"] is True


def test_too_little_text_and_empty_title_are_refused():
    with pytest.raises(HTTPException):
        ks.add_text(KDB(), ORG, "u1", "T", "short")
    with pytest.raises(HTTPException):
        ks.add_text(KDB(), ORG, "u1", "  ", "x" * 50)


def test_a_full_library_is_refused(monkeypatch):
    monkeypatch.setattr(ks, "MAX_ORG_CHUNKS", 1)
    db = KDB()
    ks.add_text(db, ORG, "u1", "One", "a " * 30)
    with pytest.raises(HTTPException) as e:
        ks.add_text(db, ORG, "u1", "Two", "b " * 30)
    assert e.value.detail["code"] == "LIBRARY_FULL"


def test_a_failed_chunk_save_leaves_nothing_behind():
    class Broken(KDB):
        def table(self, name):
            if name == "knowledge_chunks":
                raise RuntimeError("down")
            return super().table(name)
    db = Broken()
    with pytest.raises(HTTPException):
        ks.add_text(db, ORG, "u1", "T", "text " * 30)
    assert db.tables["knowledge_sources"] == []


def test_other_organisations_cannot_see_or_change_a_source():
    db = KDB()
    sid = ks.add_text(db, ORG, "u1", "Mine", "private " * 20)["id"]
    for fn in (ks.preview, ks.activate, ks.remove):
        with pytest.raises(HTTPException) as e:
            fn(db, OTHER, sid)
        assert e.value.status_code == 404
    assert ks.list_sources(db, OTHER) == []


def test_remove_deletes_the_chunks_too():
    db = KDB()
    sid = ks.add_text(db, ORG, "u1", "Gone", "text " * 30)["id"]
    ks.remove(db, ORG, sid)
    assert db.tables["knowledge_sources"] == [] and db.tables["knowledge_chunks"] == []


# -- finding the best chunks -----------------------------------------------------------------------------------
def test_query_has_only_safe_meaningful_words():
    q = ks.build_query("Hi! What's the PRICE of the wig; (delivery) -- to Abuja?? DROP TABLE x; 'or'")
    assert q is not None and all(ch.isalnum() or ch in " |:*" for ch in q)
    assert "price:*" in q and "delivery:*" in q and "abuja:*" in q
    assert "the:*" not in q and "hi:*" not in q
    assert ks.build_query("hi hello the") is None


def test_the_bot_never_gets_more_than_five_chunks():
    db = KDB()
    for i in range(12):
        ks.activate(db, ORG, ks.add_text(db, ORG, "u1", f"Doc {i}", f"price list number {i} " * 10)["id"])
    out = ks.retrieve(db, ORG, "price list", limit=50)
    assert len(out) == ks.MAX_CHUNKS and all(len(c["content"]) <= ks.MAX_CHUNK_CHARS for c in out)


def test_about_our_business_is_always_included_and_not_repeated():
    db = KDB()
    ks.upsert_text_source(db, ORG, "u1", ks.PINNED_TITLE, "What we sell: wigs and braids\n\nOpening hours: 9 to 5\n\nReply in: English")
    ks.activate(db, ORG, ks.add_text(db, ORG, "u1", "Prices", "wig price is 20000 naira " * 5)["id"])
    out = ks.retrieve(db, ORG, "how much is a wig")
    assert out[0]["title"] == ks.PINNED_TITLE and len(out) == 2
    again = ks.retrieve(db, ORG, "opening hours wigs")
    assert [c["title"] for c in again].count(ks.PINNED_TITLE) == 1
    assert [c["title"] for c in ks.retrieve(db, ORG, "hello")] == [ks.PINNED_TITLE]    # nothing matched: only the pinned chunk


def test_a_search_failure_gives_the_bot_nothing_extra_not_an_error():
    db = KDB()
    ks.activate(db, ORG, ks.add_text(db, ORG, "u1", "P", "price " * 30)["id"])
    db.rpc_error = True
    assert ks.retrieve(db, ORG, "price") == []


def test_upsert_replaces_instead_of_duplicating():
    db = KDB()
    ks.upsert_text_source(db, ORG, "u1", ks.PINNED_TITLE, "first version of the text " * 3)
    ks.upsert_text_source(db, ORG, "u1", ks.PINNED_TITLE, "second version of the text " * 3)
    assert len(db.tables["knowledge_sources"]) == 1 and "second" in db.tables["knowledge_chunks"][0]["content"]


# -- the bots use it ----------------------------------------------------------------------------------------------
def test_agent_uses_knowledge_when_present_and_old_articles_otherwise():
    from app.services import ai_agent_service as ag
    db = KDB(knowledge_base_articles=[{"org_id": ORG, "title": "Old", "content": "refund policy text", "is_published": True}])
    assert ag._fetch_kb_articles(db, ORG, "refund policy")[0]["title"] == "Old"          # no sources: unchanged behaviour
    ks.activate(db, ORG, ks.add_text(db, ORG, "u1", "New", "refund policy is seven days " * 4)["id"])
    got = ag._fetch_kb_articles(db, ORG, "refund policy")
    assert got and got[0]["chunk"] is True and got[0]["title"] == "New"


def test_prompt_gives_chunks_more_room_than_articles():
    from app.services import ai_agent_service as ag
    big = "y" * 1400
    chunks = [{"title": f"T{i}", "content": big, "chunk": True} for i in range(7)]
    articles = [{"title": f"A{i}", "content": big} for i in range(7)]
    kw = dict(org={"name": "Glow"}, ai_agent_config={}, catalog_config=None, cart=None, conversation_history=[])
    pc = ag.build_agent_system_prompt(kb_articles=chunks, **kw)
    pa = ag.build_agent_system_prompt(kb_articles=articles, **kw)
    assert pc.count("- T") == 5 and pc.count("y" * 1400) == 5
    assert pa.count("- A") == 3 and pa.count("y" * 500) == 3 and "y" * 501 not in pa


# -- the cache split -----------------------------------------------------------------------------------------------
def _capture(monkeypatch):
    from app.services import ai_service
    seen = {}

    class _Msgs:
        def create(self, **kw):
            seen.update(kw)
            return SimpleNamespace(content=[SimpleNamespace(type="text", text="ok")],
                                   usage=SimpleNamespace(input_tokens=1, output_tokens=1, cache_read_input_tokens=0, cache_creation_input_tokens=0))
    monkeypatch.setattr(ai_service, "_get_client", lambda: SimpleNamespace(messages=_Msgs()))
    monkeypatch.setattr(ai_service, "check_and_increment_token_usage", lambda *a, **k: True)
    return ai_service, seen


def test_split_prompt_caches_only_the_stable_part(monkeypatch):
    svc, seen = _capture(monkeypatch)
    system = "STABLE RULES" + svc.CACHE_SPLIT + "CHANGING KNOWLEDGE"
    svc.call_claude("hi", system=system, system_cache=True)
    blocks = seen["system"]
    assert blocks[0]["text"] == "STABLE RULES" and blocks[0]["cache_control"] == {"type": "ephemeral"}
    assert "cache_control" not in blocks[1] and "CHANGING KNOWLEDGE" in blocks[1]["text"]
    assert "<<CACHE_SPLIT>>" not in blocks[0]["text"] + blocks[1]["text"]
    # the model reads the same words in the same order as before the split existed
    assert (blocks[0]["text"] + blocks[1]["text"]) == "STABLE RULES\n\nCHANGING KNOWLEDGE\n" + svc._SECURITY_RULES


def test_without_the_marker_nothing_changes(monkeypatch):
    svc, seen = _capture(monkeypatch)
    svc.call_claude("hi", system="plain system", system_cache=True)
    assert seen["system"] == [{"type": "text", "text": "plain system\n" + svc._SECURITY_RULES, "cache_control": {"type": "ephemeral"}}]
    svc.call_claude("hi", system="plain system", system_cache=False)
    assert seen["system"] == "plain system\n" + svc._SECURITY_RULES


def test_marker_is_stripped_when_caching_is_off(monkeypatch):
    svc, seen = _capture(monkeypatch)
    svc.call_claude("hi", system="A" + svc.CACHE_SPLIT + "B", system_cache=False)
    assert seen["system"] == "A\n\nB\n" + svc._SECURITY_RULES


def test_agent_prompt_splits_before_the_knowledge_and_keeps_every_rule():
    from app.services import ai_agent_service as ag
    from app.services.ai_service import CACHE_SPLIT
    p = ag.build_agent_system_prompt(org={"name": "Glow"}, ai_agent_config={}, catalog_config=None, cart=None,
                                     conversation_history=[{"role": "user", "content": "hello"}],
                                     kb_articles=[{"title": "T", "content": "KNOWLEDGE", "chunk": True}])
    stable, _, changing = p.partition(CACHE_SPLIT)
    assert "KNOWLEDGE" not in stable and "<knowledge_base>" in changing and "<conversation_history>" in changing
    assert "CRITICAL: never invent" in changing and p.count("<<CACHE_SPLIT>>") == 1


# -- the routes -------------------------------------------------------------------------------------------------
def _client(db, template="owner", allowed=True, monkeypatch=None):
    from app.main import app
    from app.database import get_supabase
    from app.dependencies import get_current_org
    from app.services import plan_gate_service
    app.dependency_overrides[get_supabase] = lambda: db
    app.dependency_overrides[get_current_org] = lambda: {"org_id": ORG, "id": "user-1", "roles": {"template": template}}
    if monkeypatch is not None:
        monkeypatch.setattr(plan_gate_service, "allowed", lambda *a, **k: allowed)
    return TestClient(app)


@pytest.fixture(autouse=True)
def _reset():
    yield
    from app.main import app
    app.dependency_overrides.clear()


def test_upload_preview_start_and_remove_through_the_real_app(monkeypatch):
    db = KDB()
    c = _client(db, monkeypatch=monkeypatch)
    r = c.post("/api/v1/knowledge/sources/file", files={"file": ("faq.txt", b"We open at nine. " * 20, "text/plain")}, data={"title": "FAQ"})
    assert r.status_code == 200
    d = r.json()["data"]
    assert d["title"] == "FAQ" and d["status"] == "draft" and d["samples"]
    sid = d["id"]
    assert c.post(f"/api/v1/knowledge/sources/{sid}/activate").json()["data"]["status"] == "active"
    assert [s["status"] for s in c.get("/api/v1/knowledge/sources").json()["data"]] == ["active"]
    assert c.delete(f"/api/v1/knowledge/sources/{sid}").json()["data"]["removed"] is True
    assert c.get("/api/v1/knowledge/sources").json()["data"] == []
    assert [a["action"] for a in db.tables["audit_logs"]] == ["knowledge.added", "knowledge.removed"]


def test_typed_answers_route(monkeypatch):
    c = _client(KDB(), monkeypatch=monkeypatch)
    r = c.post("/api/v1/knowledge/sources/text", json={"title": "Returns", "content": "Returns are accepted within seven days. " * 3})
    assert r.status_code == 200 and r.json()["data"]["kind"] == "text"


def test_a_bad_upload_gets_a_plain_error(monkeypatch):
    c = _client(KDB(), monkeypatch=monkeypatch)
    r = c.post("/api/v1/knowledge/sources/file", files={"file": ("virus.exe", b"MZ", "application/octet-stream")})
    assert r.status_code == 422 and "PDF" in r.json()["detail"]["message"]


@pytest.mark.parametrize("template", ["sales_agent", "affiliate_partner", ""])
def test_only_owner_ops_or_admin_can_change_knowledge(monkeypatch, template):
    c = _client(KDB(), template=template, monkeypatch=monkeypatch)
    assert c.get("/api/v1/knowledge/sources").status_code == 403


def test_the_plan_must_include_knowledge(monkeypatch):
    c = _client(KDB(), allowed=False, monkeypatch=monkeypatch)
    r = c.get("/api/v1/knowledge/sources")
    assert r.status_code == 403 and r.json()["detail"]["code"] == "PLAN_REQUIRED"


def test_knowledge_routes_stay_open_to_whatsapp_only_businesses():
    import app.main as m
    from app.services import product_mode as pm
    paths = [r for r in m.app.routes if getattr(r, "path", "").startswith("/api/v1/knowledge/")]
    assert len(paths) >= 5
    assert not any(d.call is pm.crm_only for r in paths for d in r.dependant.dependencies)


def test_knowledge_feature_is_registered_and_in_the_automate_plan():
    from app.services import site_feature_registry as reg
    assert reg.FEATURES["knowledge_base"]["built"] is True
    assert "knowledge_base" in reg.DEFAULT_TIER_FEATURES["automate"]
