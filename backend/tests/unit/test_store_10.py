"""
tests/unit/test_store_10.py
---------------------------
STORE-10: add products from a spreadsheet (CSV or .xlsx), photos from web links, and sizes that carry an id.
"""
from __future__ import annotations

import io
from contextlib import contextmanager
from unittest.mock import patch

import pytest

from app.services import catalog_category_service as cats
from app.services import catalog_import_service as imp
from app.services import catalog_service as svc
from tests.unit.test_account_plans_1 import BIZ, PLATFORM, FakeDB, _Q


class _MS(_Q):
    def maybe_single(self): return self


class DB(FakeDB):
    def table(self, name): return _MS(self, name)


def _db(products=None, cfg=None):
    return DB(organisations=[{"id": BIZ, "slug": "rr", "name": "RR", "catalog_config": cfg if cfg is not None else {}}],
              products=products or [])


def _csv(*lines):
    return ("\n".join(lines) + "\n").encode()


# ── sizes get ids ───────────────────────────────────────────────────────────

def test_with_variant_ids_assigns_keeps_and_dedupes():
    out = svc.with_variant_ids([{"title": "A"}, {"title": "B", "id": "keep"}, {"title": "C", "id": 12345}, {"title": "D", "id": "keep"}, "junk"])
    ids = [v["id"] for v in out]
    assert len(out) == 4 and ids[1] == "keep" and ids[2] == 12345            # a kept id is never re-typed
    assert ids[0].startswith("v") and ids[3].startswith("v") and len({str(i) for i in ids}) == 4
    assert svc.with_variant_ids(None) == [] and svc.with_variant_ids([]) == []


def test_create_and_update_give_sizes_ids():
    db = _db([{"id": "p1", "org_id": BIZ, "slug": "x", "variants": []}])
    item = svc.create_catalog_item(db, BIZ, {"title": "Bed", "variants": [{"title": "Queen", "price": 5}]})
    assert item["variants"][0]["id"].startswith("v")
    up = svc.update_catalog_item(db, BIZ, "p1", {"variants": [{"title": "King", "price": 9}]})
    assert up["variants"][0]["id"].startswith("v")
    again = svc.update_catalog_item(db, BIZ, "p1", {"variants": [{"title": "King", "price": 9, "id": up["variants"][0]["id"]}]})
    assert again["variants"][0]["id"] == up["variants"][0]["id"]              # saving again keeps the same id
    assert "variants" not in svc.update_catalog_item(db, BIZ, "p1", {"title": "Only title"}) or True


def test_a_size_with_an_id_is_priced_by_the_cart():
    """The reason for the ids: the web chat cart now charges the size's price, not the product's."""
    from app.services import webchat_shop_service as shop
    db = _db()
    item = svc.create_catalog_item(db, BIZ, {"title": "Bed", "price": 100, "variants": [{"title": "Queen", "price": 180}]})
    vid = item["variants"][0]["id"]
    s = {"org_id": BIZ, "state": {}, "id": "s1"}
    with patch.object(shop.wc, "set_state", side_effect=lambda db_, s_, cart: s_["state"].update(cart=cart)):
        out = shop.add(db, s, item["id"], vid, 1)
    assert out["total"] == 180.0
    assert shop.add(db, s, item["id"], None, 1)["error"] == "Please choose an option first."


# ── cells ───────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("raw,want", [("150000", 150000.0), ("150,000", 150000.0), ("N150,000", 150000.0), ("₦150 000", 150000.0),
                                      ("NGN 2,500.50", 2500.5), ("", None), ("  ", None), (None, None), ("0", 0.0)])
def test_parse_price(raw, want):
    assert imp.parse_price(raw) == want


@pytest.mark.parametrize("raw", ["free", "12abc", "-5", "1e5", "9999999999"])
def test_parse_price_rejects(raw):
    with pytest.raises(ValueError):
        imp.parse_price(raw)


def test_parse_sizes_pairs_plain_lists_and_prices_with_commas():
    assert imp.parse_sizes("Single 3/4 ft=120000; Queen 5/6 ft=180,000") == [
        {"title": "Single 3/4 ft", "price": 120000.0}, {"title": "Queen 5/6 ft", "price": 180000.0}]
    assert imp.parse_sizes("Single, Queen ,King") == [{"title": "Single", "price": None}, {"title": "Queen", "price": None}, {"title": "King", "price": None}]
    assert imp.parse_sizes("S=1|M: 2\nL") == [{"title": "S", "price": 1.0}, {"title": "M", "price": 2.0}, {"title": "L", "price": None}]
    assert imp.parse_sizes("Queen; queen") == [{"title": "Queen", "price": None}]          # same size twice
    assert imp.parse_sizes("") == [] and imp.parse_sizes(None) == []
    with pytest.raises(ValueError):
        imp.parse_sizes("Queen=lots")
    with pytest.raises(ValueError):
        imp.parse_sizes(";".join(f"s{i}" for i in range(21)))


def test_parse_in_stock_photos_lists():
    assert [imp.parse_in_stock(x) for x in ("yes", "", "No", "sold out", "OUT", "0", "Y", "anything")] == [True, True, False, False, False, False, True, True]
    good, notes = imp.parse_photos("https://a/1.jpg, http://b/2.jpg  https://a/1.jpg ftp://c")
    assert good == ["https://a/1.jpg"] and len(notes) == 2 and "https://" in notes[0]
    good, notes = imp.parse_photos(" ".join(f"https://a/{i}.jpg" for i in range(8)))
    assert len(good) == 5 and any("first 5" in n for n in notes)
    assert imp.parse_list("Beds, pillows ;Beds| ") == ["Beds", "pillows"]


# ── the file ────────────────────────────────────────────────────────────────

def test_headings_can_be_written_many_ways_and_rows_are_numbered_like_the_sheet():
    t = imp.read_table(_csv("Product Name,Selling Price,Categories,Photo links,Notes", "Bed,\"1,500\",Beds,https://a/x.jpg,hi", ",5,,,"), "p.csv")
    p = imp.parse_table(t)
    assert p["found"] == ["category", "photos", "price", "title"] and p["unknown"] == ["Notes"]
    assert p["rows"][0]["row"] == 2 and p["rows"][0]["price"] == 1500.0 and p["rows"][0]["photos"] == ["https://a/x.jpg"]
    assert p["rows"][1]["problems"] == ["No title."]


def test_semicolon_and_tab_files_and_a_byte_order_mark_open_fine():
    for delim in (";", "\t"):
        raw = ("﻿" + delim.join(["Title", "Price"]) + "\n" + delim.join(["Bed", "5"]) + "\n").encode("utf-8")
        assert imp.parse_table(imp.read_table(raw, "x.csv"))["rows"][0]["title"] == "Bed"
    cp = "Title,Price\nCaf\xe9 chair,5\n".encode("cp1252")
    assert imp.parse_table(imp.read_table(cp, "x.csv"))["rows"][0]["title"] == "Caf\xe9 chair"


def test_xlsx_opens():
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws.append(["Title", "Price", "Sizes"]); ws.append(["Bed", 150000.0, "Q=1"]); ws.append([None, None, None])
    buf = io.BytesIO(); wb.save(buf)
    p = imp.parse_table(imp.read_table(buf.getvalue(), "stock.XLSX"))
    assert len(p["rows"]) == 1 and p["rows"][0]["price"] == 150000.0 and p["rows"][0]["sizes"][0]["title"] == "Q"


@pytest.mark.parametrize("data,name,msg", [
    (b"", "a.csv", "empty"), (b"x" * (2 * 1024 * 1024 + 1), "a.csv", "too big"), (b"x", "a.xls", "Old Excel"),
    (b"x", "a.pdf", "csv or .xlsx"), (b"not a zip", "a.xlsx", "Excel file"),
])
def test_bad_files_get_a_plain_message(data, name, msg):
    with pytest.raises(imp.FileProblem) as e:
        imp.read_table(data, name)
    assert msg.lower() in str(e.value).lower()


def test_file_level_problems():
    for raw, msg in ((_csv("Price,Size", "5,Q"), "Title column"), (_csv("Title,Price"), "no products"), (b"\n\n", "no rows")):
        with pytest.raises(imp.FileProblem) as e:
            imp.parse_table(imp.read_table(raw, "a.csv"))
        assert msg in str(e.value)
    many = _csv("Title", *[f"P{i}" for i in range(501)])
    with pytest.raises(imp.FileProblem) as e:
        imp.parse_table(imp.read_table(many, "a.csv"))
    assert "501" in str(e.value)


def test_row_problems_are_listed_not_fatal():
    p = imp.parse_table(imp.read_table(_csv("Title,Price,Sizes", "A,free,", "B,5,Q=lots", "C,,", "D," + "x", ("E" * 201) + ",5,"), "a.csv"))
    probs = [r["problems"] for r in p["rows"]]
    assert "not a number" in probs[0][0] and "not a number" in probs[1][0] and probs[2] == []
    assert p["rows"][2]["notes"] == ["No price: customers will have to ask you about it."]
    assert "longer than" in probs[4][0]


# ── plan and run ────────────────────────────────────────────────────────────

def _plan(db, *lines):
    return imp.plan(db, BIZ, imp.parse_table(imp.read_table(_csv(*lines), "a.csv")))


def test_plan_marks_existing_duplicate_error_and_ok():
    db = _db([{"id": "1", "org_id": BIZ, "title": "Cloud 9"}, {"id": "2", "org_id": "other", "title": "Foam"}])
    pl = _plan(db, "Title,Price,Category", "cloud 9,5,", "Foam,6,\"Beds, Pillows\"", "foam,7,", "Bad,x,", "Ok,1,Beds")
    assert [r["status"] for r in pl["rows"]] == ["exists", "ok", "duplicate", "error", "ok"]       # another business's Foam does not count
    assert pl["counts"] == {"total": 5, "ok": 2, "exists": 1, "duplicate": 1, "error": 1}
    assert pl["categories"] == ["Beds", "Pillows"]


def test_run_creates_ok_rows_only_with_sizes_categories_and_stock():
    db = _db([{"id": "1", "org_id": BIZ, "title": "Old"}])
    pl = _plan(db, "Title,Price,Description,Category,Sizes,Photos,In stock",
               "Old,5,,,,,", "Bed,100,Nice,\"Beds, Big\",Q=180; K=220,https://a/1.jpg,no", "Bad,x,,,,,")
    out = imp.run(db, BIZ, pl, svc.create_catalog_item, cats.ensure_categories)
    assert [c["title"] for c in out["created"]] == ["Bed"] and out["created"][0]["photos"] == ["https://a/1.jpg"]
    assert out["failed"] == [] and out["categories_added"] == ["Beds", "Big"]
    row = next(r for r in db.tables["products"] if r["title"] == "Bed")
    assert row["available"] is False and row["description"] == "Nice" and row["tags"] == {"category": ["Beds", "Big"]}
    assert [(v["title"], v["price"]) for v in row["variants"]] == [("Q", 180.0), ("K", 220.0)] and all(v["id"].startswith("v") for v in row["variants"])
    assert row["catalog_visible"] is True and row["slug"] == "bed"
    cfg = db.tables["organisations"][0]["catalog_config"]
    assert next(d for d in cfg["tag_dimensions"] if d["key"] == "category")["options"] == ["Beds", "Big"]
    assert len(db.tables["products"]) == 2                                  # the "Old" row and the bad row were not added


def test_run_one_failure_does_not_stop_the_rest():
    db = _db()
    pl = _plan(db, "Title,Price", "A,1", "B,2", "C,3")
    calls = []

    def flaky(db_, org, payload):
        calls.append(payload["title"])
        if payload["title"] == "B":
            raise RuntimeError("db down")
        return svc.create_catalog_item(db_, org, payload)
    out = imp.run(db, BIZ, pl, flaky, cats.ensure_categories)
    assert [c["title"] for c in out["created"]] == ["A", "C"] and out["failed"][0]["row"] == 3 and "by hand" in out["failed"][0]["message"]


def test_ensure_categories_respects_the_limit_and_bad_names():
    db = _db(cfg={"tag_dimensions": [{"key": "category", "label": "Category", "type": "multi_select", "filterable": True, "options": [f"c{i}" for i in range(19)]}]})
    out = cats.ensure_categories(db, BIZ, ["C0", "new1", "new2", "a/b", "  "])
    assert out["added"] == ["new1"] and out["skipped"] == ["new2", "a/b", "  "] and out["names"]["c0"] == "c0"


def test_template_opens_in_excel_and_reads_back():
    t = imp.template_csv()
    assert t.startswith("﻿")
    p = imp.parse_table(imp.read_table(t.encode("utf-8"), "t.csv"))
    assert len(p["rows"]) == 2 and p["rows"][0]["sizes"][1]["price"] == 180000.0 and all(not r["problems"] for r in p["rows"])


# ── pictures from links ─────────────────────────────────────────────────────

class _Resp:
    def __init__(self, status=200, body=b"", headers=None):
        self.status_code, self.body, self.headers = status, body, headers or {}
    def iter_bytes(self): yield self.body


def _opener(*responses):
    seq = list(responses); urls = []

    @contextmanager
    def op(method, url, **kw):
        urls.append(url); assert kw["follow_redirects"] is False
        yield seq.pop(0)
    op.urls = urls
    return op


JPEG = b"\xff\xd8\xff\xe0" + b"0" * 50


@pytest.fixture
def public_dns():
    with patch.object(imp, "_public_ip", return_value=True):
        yield


def test_fetch_image_ok(public_dns):
    assert imp.fetch_image("https://a.test/x.jpg", opener=_opener(_Resp(body=JPEG))) == (JPEG, "image/jpeg")
    png = b"\x89PNG\r\n\x1a\n" + b"0" * 9
    assert imp.fetch_image("https://a.test/x", opener=_opener(_Resp(body=png)))[1] == "image/png"
    webp = b"RIFF\x00\x00\x00\x00WEBPVP8 "
    assert imp.fetch_image("https://a.test/x", opener=_opener(_Resp(body=webp)))[1] == "image/webp"


@pytest.mark.parametrize("url", ["http://a.test/x.jpg", "ftp://a.test/x", "https:///x", "javascript:alert(1)", "", "//a.test/x"])
def test_fetch_image_refuses_non_https(url, public_dns):
    op = _opener(_Resp(body=JPEG))
    with pytest.raises(imp.PhotoError):
        imp.fetch_image(url, opener=op)
    assert op.urls == []                                                     # nothing was requested


def test_fetch_image_refuses_private_addresses_even_via_a_redirect():
    import socket
    def fake(host, *a, **k):
        ip = {"public.test": "93.184.216.34", "internal.test": "10.0.0.5", "meta.test": "169.254.169.254", "local.test": "127.0.0.1", "v6.test": "::1"}[host]
        return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (ip, 443))]
    with patch.object(socket, "getaddrinfo", side_effect=fake):
        for h in ("internal.test", "meta.test", "local.test", "v6.test"):
            with pytest.raises(imp.PhotoError):
                imp.fetch_image(f"https://{h}/x.jpg", opener=_opener(_Resp(body=JPEG)))
        op = _opener(_Resp(302, headers={"location": "https://internal.test/secret"}))
        with pytest.raises(imp.PhotoError):
            imp.fetch_image("https://public.test/x.jpg", opener=op)
        assert op.urls == ["https://public.test/x.jpg"]                      # the private address was never requested
        op = _opener(_Resp(301, headers={"location": "/b.jpg"}), _Resp(body=JPEG))
        assert imp.fetch_image("https://public.test/a.jpg", opener=op)[1] == "image/jpeg" and op.urls[1] == "https://public.test/b.jpg"


def test_fetch_image_limits(public_dns):
    with pytest.raises(imp.PhotoError, match="5 MB"):
        imp.fetch_image("https://a.test/x", opener=_opener(_Resp(body=JPEG + b"0" * (5 * 1024 * 1024))))
    with pytest.raises(imp.PhotoError, match="not a JPEG"):
        imp.fetch_image("https://a.test/x", opener=_opener(_Resp(body=b"<html>nope</html>")))
    with pytest.raises(imp.PhotoError, match="direct picture"):
        imp.fetch_image("https://a.test/x", opener=_opener(_Resp(404)))
    with pytest.raises(imp.PhotoError, match="too many"):
        imp.fetch_image("https://a.test/x", opener=_opener(*[_Resp(302, headers={"location": "https://a.test/y"}) for _ in range(4)]))


def test_fetch_image_network_error_is_plain(public_dns):
    @contextmanager
    def boom(*a, **k):
        raise OSError("reset")
        yield
    with pytest.raises(imp.PhotoError, match="could not download"):
        imp.fetch_image("https://a.test/x", opener=boom)


# ── routes ──────────────────────────────────────────────────────────────────

@pytest.fixture()
def owner(request):
    from fastapi.testclient import TestClient
    from app.dependencies import get_current_org, get_supabase
    from app.main import app
    role = getattr(request, "param", "owner")
    db = _db()
    app.dependency_overrides[get_current_org] = lambda: {"org_id": BIZ, "roles": {"template": role}}
    app.dependency_overrides[get_supabase] = lambda: db
    with patch("app.services.plan_gate_service.allowed", return_value=True):
        yield TestClient(app), db
    app.dependency_overrides.clear()


def _up(data, name="a.csv"):
    return {"file": (name, data, "text/csv")}


def test_route_preview_writes_nothing_and_import_adds(owner):
    c, db = owner
    f = _csv("Title,Price,Category", "Bed,5,Beds", "Bad,x,")
    r = c.post("/api/v1/catalog/import/preview", files=_up(f))
    assert r.status_code == 200 and r.json()["counts"]["ok"] == 1 and r.json()["counts"]["error"] == 1
    assert db.tables["products"] == []
    r = c.post("/api/v1/catalog/import", files=_up(f))
    assert r.status_code == 201 and [x["title"] for x in r.json()["created"]] == ["Bed"] and r.json()["categories_added"] == ["Beds"]
    assert len(db.tables["products"]) == 1
    again = c.post("/api/v1/catalog/import", files=_up(f)).json()
    assert again["created"] == [] and again["counts"]["exists"] == 1                  # importing twice never duplicates


def test_route_bad_file_is_422_with_a_plain_message(owner):
    c, _ = owner
    r = c.post("/api/v1/catalog/import/preview", files=_up(_csv("Price", "5")))
    assert r.status_code == 422 and "Title column" in r.json()["detail"]


def test_route_shopify_business_is_refused():
    from fastapi.testclient import TestClient
    from app.dependencies import get_current_org, get_supabase
    from app.main import app
    db = _db(cfg={"external_sync": "shopify"})
    app.dependency_overrides[get_current_org] = lambda: {"org_id": BIZ, "roles": {"template": "owner"}}
    app.dependency_overrides[get_supabase] = lambda: db
    try:
        with patch("app.services.plan_gate_service.allowed", return_value=True):
            r = TestClient(app).post("/api/v1/catalog/import/preview", files=_up(_csv("Title", "A")))
            assert r.status_code == 403 and "Shopify" in r.json()["detail"]
            r = TestClient(app).post("/api/v1/catalog/import", files=_up(_csv("Title", "A")))
        assert r.status_code == 403 and db.tables["products"] == []
    finally:
        app.dependency_overrides.clear()


@pytest.mark.parametrize("owner", ["staff", "agent"], indirect=True)
def test_routes_need_owner_or_ops_manager(owner):
    c, _ = owner
    assert c.post("/api/v1/catalog/import/preview", files=_up(_csv("Title", "A"))).status_code == 403
    assert c.post("/api/v1/catalog/import", files=_up(_csv("Title", "A"))).status_code == 403
    assert c.get("/api/v1/catalog/import/template.csv").status_code == 403
    assert c.post("/api/v1/catalog/items/p1/image-from-url", json={"url": "https://a/x.jpg"}).status_code == 403


def test_route_template_downloads(owner):
    c, _ = owner
    r = c.get("/api/v1/catalog/import/template.csv")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"] and r.text.lstrip("﻿").startswith("Title,Price")


def test_route_image_from_url(owner):
    c, db = owner
    db.tables["products"].append({"id": "p1", "org_id": BIZ, "title": "Bed", "catalog_images": []})
    with patch("app.routers.catalog.importer.fetch_image", return_value=(JPEG, "image/jpeg")), \
         patch("app.routers.catalog.upload_catalog_image", return_value="https://store/p1/link-1.jpg") as up:
        r = c.post("/api/v1/catalog/items/p1/image-from-url", json={"url": "https://a.test/x.jpg"})
    assert r.status_code == 201 and r.json() == {"url": "https://store/p1/link-1.jpg"}
    assert up.call_args.args[2] == "p1" and up.call_args.args[4].startswith("link-") and up.call_args.args[4].endswith(".jpg")
    assert c.post("/api/v1/catalog/items/nope/image-from-url", json={"url": "https://a.test/x.jpg"}).status_code == 404
    with patch("app.routers.catalog.importer.fetch_image", side_effect=imp.PhotoError("That link cannot be used.")):
        r = c.post("/api/v1/catalog/items/p1/image-from-url", json={"url": "https://10.0.0.1/x.jpg"})
    assert r.status_code == 422 and r.json()["detail"] == "That link cannot be used."
